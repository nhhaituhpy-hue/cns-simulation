#!/usr/bin/env python3
"""Build six editable default-reference DOCX files from the native form export.

Requires python-docx >= 1.2. Run the companion export-ai-authoring-forms.mts
first (its header documents the esbuild command), then:
  python scripts/generate-ai-authoring-templates.py --input forms.json --output public/templates

No ranges, defaults, or field inventories are maintained in this generator.
The input is an allowlisted projection of form(createDefault()), not a scenario.
"""
from __future__ import annotations
import argparse
from collections import OrderedDict
from datetime import datetime, timezone
import json
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

MODULES = {"dvor-1150", "dvor-1150a", "dme-1119a", "dvor-220", "dme-320", "ads-b"}
FIXED_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)


def clean_package(document):
    """Keep used-style dependencies; remove unsafe/unused python-docx parts."""
    for rid, rel in list(document.part.rels.items()):
        if rel.reltype.endswith(("/customXml", "/stylesWithEffects")):
            document.part.drop_rel(rid)
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    used = {p.style.style_id for p in paragraphs} | {"Normal", "DefaultParagraphFont", "TableNormal", "TableGrid", "NoList"}
    for element in document.part.numbering_part.element.iter(qn("w:pStyle")):
        used.add(element.get(qn("w:val")))
    styles = document.styles.element
    by_id = {e.get(qn("w:styleId")): e for e in styles.findall(qn("w:style"))}
    while True:
        previous = set(used)
        for key in previous:
            element = by_id.get(key)
            if element is None:
                continue
            for tag in ("basedOn", "next", "link"):
                ref = element.find(qn("w:" + tag))
                if ref is not None:
                    used.add(ref.get(qn("w:val")))
        if used == previous:
            break
    for key, element in by_id.items():
        if key not in used:
            styles.remove(element)
    latent = styles.find(qn("w:latentStyles"))
    if latent is not None:
        styles.remove(latent)


def friendly_group(value):
    value = re.sub(r"([a-z])([A-Z])", r"\1 \2", value)
    value = re.sub(r"\b(tx|mon)([12])\b", lambda m: m[1].upper() + m[2], value)
    return value[:1].upper() + value[1:]


def limits(schema):
    parts = []
    if schema.get("unit"):
        parts.append("Đơn vị: " + schema["unit"])
    if schema.get("options"):
        parts.append("Chọn: " + "; ".join(option["label"] for option in schema["options"]))
    elif schema["type"] == "boolean":
        parts.append("Bật / Tắt")
    for key, prefix in (("min", "≥ "), ("max", "≤ "), ("step", "Bước "), ("maxLength", "Tối đa ký tự: "), ("maxItems", "Tối đa mục: ")):
        if key in schema:
            parts.append(prefix + str(schema[key]))
    if schema.get("nullable"):
        parts.append("Có thể chưa đặt")
    if schema["type"] == "number" and "min" not in schema and "max" not in schema:
        parts.append("Chưa công bố khoảng số trong biểu mẫu")
    return " · ".join(parts)


def display(value, schema):
    if value is None:
        return "Chưa đặt"
    if isinstance(value, bool):
        return "Bật" if value else "Tắt"
    for option in schema.get("options", []):
        if option["value"] == value:
            return option["label"]
    if value == "":
        return "Trống"
    return str(value)


def rows(schema, value, prefix=""):
    """Expand structured controls, preserving labels and actual default entries."""
    label = (prefix + " / " if prefix else "") + schema["label"]
    if schema["type"] == "object":
        yield label, limits(schema), "Nhóm thông số"
        data = value if isinstance(value, dict) else {}
        for key, child in schema.get("properties", {}).items():
            yield from rows(child, data.get(key), label)
    elif schema["type"] == "array":
        yield label, limits(schema), "Danh sách: " + str(len(value) if isinstance(value, list) else 0) + " mục"
        items = value if isinstance(value, list) else []
        if schema.get("item"):
            if items:
                for index, item in enumerate(items, 1):
                    yield from rows(schema["item"], item, label + " — mục " + str(index))
            else:
                # Blank item is a schema reference, never an invented baseline value.
                yield from rows(schema["item"], None, label + " — mục trống (không phải dữ liệu)")
        for variant in schema.get("variants", []):
            yield from rows(variant["schema"], None, label + " — " + variant["label"])
    else:
        yield label, limits(schema), display(value, schema)


def table(document, entries):
    result = document.add_table(rows=1, cols=3)
    result.style = "Table Grid"
    result.autofit = False
    widths = [3.8, 1.15, 1.9]
    for cell, text, width in zip(result.rows[0].cells, ["Thông số / đơn vị / giới hạn", "Mặc định mô phỏng", "Giá trị tài liệu / Nguồn/trang"], widths):
        cell.text = text
        cell.width = Inches(width)
        cell.paragraphs[0].runs[0].bold = True
    header = OxmlElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(header)
    count = 0
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    for label, constraint, value in entries:
        cells = result.add_row().cells
        # Widths belong to the table grid. Repeating cell properties on every
        # row consumes the production XML token/node budget on large forms.
        for cell in cells:
            properties = cell._tc.find(qn("w:tcPr"))
            if properties is not None:
                cell._tc.remove(properties)
        cells[0].text = label + (" · " + constraint if constraint else "")
        cells[1].text = value
        cells[2].text = "Giá trị: ______; Nguồn/trang: ______"
        count += 1
    return count


def workflow(document, module):
    document.add_heading("C. Tham khảo các bước tiếp theo — không bắt buộc điền trong Word", 1)
    document.add_paragraph("Tên, mô tả, độ khó và các mục sau có thể điền trong modal. Đây chỉ là ghi chú tùy chọn, không phải yêu cầu phải điền lỗi hay đáp án trong tài liệu nền.")
    document.add_paragraph("1. Trạng thái khởi đầu: xem các giá trị nền đã chuẩn hóa; kiểm tra nguồn và xác nhận từng nhóm trong trình hướng dẫn.")
    document.add_paragraph("2. Lỗi đưa vào: chọn lỗi hoặc yêu cầu tình huống mà đúng thiết bị hỗ trợ trong modal; không dùng cột thông số nền để ghi giá trị gây lỗi.")
    document.add_paragraph("3. Điều kiện đạt: đặt mục tiêu bài tập trong modal; không thay ngưỡng bảo vệ để né lỗi.")
    if module["moduleId"] in {"dvor-1150", "dvor-1150a", "dme-1119a"}:
        document.add_paragraph("4. Chẩn đoán hai bước: khi bài tập cần, dùng các màn hình, thao tác và phần cứng được trình hướng dẫn liệt kê; không tự đặt vị trí phần cứng.")
        document.add_paragraph("5. Quyền chỉnh sửa của thí sinh: chọn các trường được phép sửa trong modal, rồi xem lại trước xác nhận cuối.")
    elif module["moduleId"] in {"dvor-220", "dme-320"}:
        document.add_paragraph("4. Thao tác và quan sát: dùng các mục RF, Monitor, Plant, lỗi có kiểu và Forced Monitor của đúng thiết bị. Không có bước chẩn đoán hai bước kiểu SELEX trong biểu mẫu này.")
        document.add_paragraph("5. Kiểm tra phạm vi: chính sách đăng nhập/quyền sửa được bảo vệ theo mô hình; không tạo mục quyền sửa hoặc đáp án chẩn đoán không được hỗ trợ.")
    else:
        document.add_paragraph("4. Thao tác chuẩn và phần cứng: dùng vai trò terminal, menu và phần cứng được modal hỗ trợ. Replay thao tác không chứng minh GPS đồng bộ hay cảm biến ngoài thực tế hồi phục.")
        document.add_paragraph("5. Kiểm tra phạm vi: xác nhận Site, cảm biến và vai trò; không cung cấp thông tin đăng nhập hoặc tự đặt quyền sửa ngoài mô hình.")
    document.add_paragraph("Các nhóm thật trong trình hướng dẫn: " + "; ".join(s["title"] for s in module["sections"]) + ".")
    document.add_paragraph("Ghi chú tình huống tùy chọn: ______. Thông tin còn thiếu hoặc mâu thuẫn cần hỏi lại: ______.")
    document.add_paragraph("Chỉ xác nhận cuối mới tạo kịch bản chưa xuất bản. Giám khảo xem lại trong simulator trước khi đưa vào Ôn tập/Kỳ thi; mẫu này không tự xuất bản, không chấm điểm và không xác minh khôi phục.")


def generate(module, output):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.6)
    section.left_margin = section.right_margin = Inches(.65)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor.from_string("17365D")
    document.core_properties.author = "CNS Simulation Lab"
    document.core_properties.title = "Mẫu thông số nền — " + module["label"]
    document.core_properties.created = document.core_properties.modified = FIXED_TIME
    document.add_heading(document.core_properties.title, 0)
    document.add_paragraph("Tài liệu tham chiếu mặc định • Chỉnh sửa trực tiếp bằng Word", style="Subtitle")
    document.add_heading("A. Cách dùng và nguồn", 1)
    for text in [
        "Bạn không phải viết toàn bộ kịch bản trong Word. Tài liệu này chỉ cung cấp thông số nền/tham chiếu cho đúng thiết bị " + module["label"] + ". Có thể tải manual có lớp chữ thay cho mẫu này.",
        "Điền giá trị có nguồn trong cột Giá trị tài liệu, kèm tên tài liệu/phiên bản và Nguồn/trang. Để trống khi chưa biết hoặc không áp dụng; ô trống không có nghĩa là 0, không tắt tính năng và không yêu cầu AI tự đoán.",
        "Mặc định mô phỏng là giá trị từ biểu mẫu native, chỉ để đối chiếu; không phải yêu cầu sửa và không phải thông số danh định được nhà sản xuất xác nhận. Giới hạn/lựa chọn lấy từ cùng biểu mẫu; nơi chưa có khoảng số phải đối chiếu nguồn, không tự đặt giới hạn.",
        "Tải tệp lên modal để AI chuẩn hóa giá trị nền, kiểm tra nguồn/thiếu dữ liệu/mâu thuẫn và đề xuất thay đổi. Sau đó xem các trường đã điền, sửa nếu cần và xác nhận từng bước trong trình hướng dẫn. Tải lên không tự tạo kịch bản hoặc đổi mặc định chung.",
        "Chỉ cột Giá trị tài liệu do bạn điền mới là dữ liệu đề xuất từ tài liệu. Hướng dẫn, dòng trống, ví dụ và cột mặc định là nội dung tham khảo, không phải dữ liệu lỗi/khôi phục.",
        "Mẫu đào tạo này không phải manual nhà sản xuất. Không đưa bí mật, mật khẩu, thông tin đăng nhập hoặc bài làm học viên vào tệp. Chỉ dùng số liệu mô phỏng hoặc tài liệu được phép chia sẻ.",
        "Nguồn chung (tên/phiên bản tài liệu): ______. Trang/đoạn: ______. Nếu chưa có: ghi Chưa cung cấp nguồn kỹ thuật. Nguồn riêng ở từng dòng được ưu tiên.",
        "Ví dụ cách điền — minh họa cách ghi nguồn, không phải yêu cầu: chép một giá trị đã kiểm chứng vào cột Giá trị tài liệu, rồi ghi tên tài liệu và trang tương ứng; không sao chép ví dụ làm số liệu của thiết bị.",
    ]:
        document.add_paragraph(text)
    document.add_heading("B. Thông số nền có thể chỉnh sửa", 1)
    document.add_paragraph("Mỗi dòng giữ tên trường của biểu mẫu. Giữ đúng đơn vị; không tự chuyển đổi. Khi một nhóm có nhiều máy/kênh, đối chiếu đúng TX/MON hoặc cảm biến trước khi xác nhận.")
    groups = OrderedDict()
    for field in module["fields"]:
        group = friendly_group(field["group"])
        if module["moduleId"] == "ads-b":
            match = re.match(r"sites\.(\d+)(?:\.sensor([AB]))?", field["id"])
            if match:
                group = "Site " + str(int(match[1]) + 1) + (" — Cảm biến " + match[2] if match[2] else "") + " — " + group
        groups.setdefault((field["sectionId"], group), []).append(field)
    row_count = 0
    for (_, group), fields in groups.items():
        document.add_heading(group, 2)
        row_count += table(document, (row for field in fields for row in rows(field["schema"], field["value"])))
    workflow(document, module)
    clean_package(document)
    path = output / ("ai-scenario-" + module["moduleId"] + ".docx")
    document.save(path)
    # Stabilize ZIP timestamps for byte-identical regeneration; do not rewrite XML.
    with ZipFile(path) as archive:
        parts = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for entry, data in parts:
            entry.date_time = (2026, 1, 1, 0, 0, 0)
            archive.writestr(entry, data)
    print(json.dumps({"moduleId": module["moduleId"], "baselineFields": len(module["fields"]), "rows": row_count, "bytes": path.stat().st_size, "path": str(path)}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="JSON exported from the real native forms")
    parser.add_argument("--output", required=True, type=Path, help="Destination for the six DOCX files")
    args = parser.parse_args()
    manifest = json.loads(args.input.read_text(encoding="utf-8"))
    modules = manifest["modules"]
    if manifest.get("version") != 1 or {m["moduleId"] for m in modules} != MODULES or len(modules) != len(MODULES):
        raise ValueError("Expected exactly the six supported native modules")
    args.output.mkdir(parents=True, exist_ok=True)
    for module in modules:
        generate(module, args.output)


if __name__ == "__main__":
    main()
